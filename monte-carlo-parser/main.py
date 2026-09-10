from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from typing import Optional
import json
import os
import sys
# For Download the Results of Excel
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from io import BytesIO

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.drawing.image import Image as XLImage

sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from src.etl.extract import extract
from src.etl.transform import transform
from src.etl.load import load
from src.config.settings import load_config
from src.utils.validators import validate_simulation_input
from src.utils.das_monte_carlo import generate_das_report
from src.utils.schemas import DASSimulationRequest, SimulationRequest
from src.utils.flatten_metadata import _flatten_metadata

app = FastAPI(title="Monte Carlo Simulation API")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
static_dir = os.path.join(BASE_DIR, "data/processed")
os.makedirs(static_dir, exist_ok=True)
app.mount('/static', StaticFiles(directory='data/processed'), name='static')

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/")
async def root():
    return {"message": "Monte Carlo Simulation API is running"}

@app.post("/run-simulation")
async def run_simulation(request: SimulationRequest):
    try:
        validate_simulation_input(request.data)

        config = load_config()
        expected_months = set(config.columns.months)
        for year, month_vals in request.data.items():
            if set(month_vals.keys()) != expected_months:
                raise HTTPException(
                    status_code=400,
                    detail=f"Year {year} must have exactly these months: {sorted(expected_months)}"
                )

        df_hist = extract(manual_data=request.data)

        forecast_years = request.forecast_years if request.forecast_years is not None else config.monte_carlo.forecast_years

        result = transform(
            df_hist=df_hist,
            forecast_years=forecast_years
        )
        if result is None:
            raise HTTPException(status_code=400, detail="Transform failed")

        load(result=result)

        return {
            "status": "success",
            "message": "Simulation completed Successfully"
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/upload-excel")
async def run_from_excel(
    file: UploadFile = File(...),
    forecast_years: Optional[int] = None
):
    config = load_config()
    raw_dir = config.data.raw_dir
    os.makedirs(raw_dir, exist_ok=True)

    file_path = os.path.join(raw_dir, file.filename)
    with open(file_path, "wb") as f:
        f.write(await file.read())

    try:
        df_hist = extract(file_path=file_path)

        result = transform(df_hist=df_hist, forecast_years=forecast_years)
        if result is None:
            raise HTTPException(400, "Transform failed")
        load(result=result)

        return {
            "status": "success", "message": "Simulation from Excel completed successfully"
        }
    
    except Exception as e:
        raise HTTPException(500, "Failed, Internal Server Error")

@app.get("/get-data")
async def getAllData():
    try:
        config = load_config()

        csv_path = config.data.processed_path
        out_dir = os.path.dirname(os.path.abspath(csv_path))
        metrics_path = os.path.join(out_dir, 'metrics.json')
        metadata_path = os.path.join(out_dir, 'metadata.json')

        if not os.path.exists(metrics_path) or not os.path.exists(metadata_path):
            raise HTTPException(404, "No simulation results found. Run a simulation first.")
        
        with open(metrics_path, 'r') as f:
            metrics = json.load(f)
        with open(metadata_path, 'r') as f:
            metadata = json.load(f)
        

        return {
            "status": "success",
            "metrics": metrics,
            "plot_urls": {
                "timeseries": "/static/timeseries_plot.png",
                "heatmap": "/static/heatmap_plot.png",
                "das_chart": "/static/grafik_konversi_das.png"
            },
            "forecast_years_used": metadata.get("forecast_years"),
            "random_seed_used": metadata.get("random_seed"),
            "das_parameters": metadata.get("das_parameters")
        }
    except Exception as e:
        raise HTTPException(500, detail=str(e))

@app.post("/post-das")
async def postDasParameter(req: DASSimulationRequest):
    try:
        config = load_config()
        processed_dir = os.path.dirname(os.path.abspath(config.data.processed_path))
        
        # Path file data simulasi Monte Carlo & metadata
        pred_csv_path = os.path.join(processed_dir, 'das_monte_carlo.csv')
        metadata_path = os.path.join(processed_dir, 'metadata.json')
        das_excel_path = os.path.join(processed_dir, 'Konversi_Curah_Hujan_DAS.xlsx')
        das_chart_path = os.path.join(processed_dir, 'grafik_konversi_das.png')

        # 1. Pastikan data simulasi Monte Carlo sudah pernah dihitung/tersimpan
        if not os.path.exists(pred_csv_path):
            raise HTTPException(
                status_code=400, 
                detail="Data simulasi Monte Carlo belum ditemukan. Jalankan simulasi Monte Carlo terlebih dahulu."
            )

        # 2. Baca data hasil prediksi Monte Carlo yang sudah ada
        df_predictions = pd.read_csv(pred_csv_path, index_col=0)

        # 3. Jalankan ulang kalkulasi DAS menggunakan parameter BARU dari req
        das_hasil = generate_das_report(
            data=df_predictions,
            cn_value=req.cn_value if req.cn_value is not None else 75.0,
            area_km2=req.area_km2 if req.area_km2 is not None else 100.0,
            n_trials=req.n_trials if req.n_trials is not None else 500,
            output_excel=das_excel_path,
            output_chart=das_chart_path
        )

        # 4. Baca metadata.json yang lama (jika ada)
        metadata_dict = {}
        if os.path.exists(metadata_path):
            with open(metadata_path, 'r') as f:
                try:
                    metadata_dict = json.load(f)
                except json.JSONDecodeError:
                    metadata_dict = {}

        # 5. Perbarui bagian das_parameters di metadata.json
        metadata_dict['das_parameters'] = das_hasil.get('das_parameters')

        # 6. Simpan kembali metadata.json yang sudah diperbarui
        with open(metadata_path, 'w') as f:
            json.dump(metadata_dict, f, indent=4)

        return {
            "status": "success",
            "message": "Parameter DAS berhasil diperbarui dan dihitung ulang!"
        }

    except HTTPException as he:
        raise he
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/export-excel")
async def export_excel():
    try: 
        config = load_config()
        csv_path = config.data.processed_path
        out_dir = os.path.dirname(os.path.abspath(csv_path))
        metrics_path = os.path.join(out_dir, 'metrics.json')
        metadata_path = os.path.join(out_dir, 'metadata.json')

        if not os.path.exists(csv_path) or not os.path.exists(metrics_path):
            raise HTTPException(404, "No simulation results found. Run a simulation first.")

        with open(metrics_path, 'r') as f:
            metrics = json.load(f)
        with open(metadata_path, 'r') as f:
            metadata = json.load(f)

        # same df_combined that produced the seaborn heatmap
        index_col = config.columns.index  # "Tahun"
        df_combined = pd.read_csv(csv_path, index_col=0)

        wb = Workbook()

        # Sheet 1: metadata
        ws_meta = wb.active
        ws_meta.title = "Metadata"
        ws_meta.append(["Key", "Value"])
        for k, v in _flatten_metadata(metadata):
            ws_meta.append([k, v])

        # Sheet 2: metrics
        ws_metrics = wb.create_sheet("Metrics")
        if isinstance(metrics, dict):
            ws_metrics.append(["Metric", "Value"])
            for k, v in metrics.items():
                ws_metrics.append([k, v])

        # Sheet 3: real heatmap data w/ conditional formatting (matches YlGnBu)
        ws_heat = wb.create_sheet("Heatmap Data")
        ws_heat.append([index_col] + list(df_combined.columns))
        for year, row in df_combined.iterrows():
            ws_heat.append([year] + [round(v, 1) for v in row.values])

        n_rows, n_cols = df_combined.shape
        first_cell = ws_heat.cell(row=2, column=2).coordinate
        last_cell = ws_heat.cell(row=1 + n_rows, column=1 + n_cols).coordinate

        rule = ColorScaleRule(
            start_type="min", start_color="FFFFE5",
            mid_type="percentile", mid_value=50, mid_color="7FCDBB",
            end_type="max", end_color="253494",
        )
        ws_heat.conditional_formatting.add(f"{first_cell}:{last_cell}", rule)

        # format numbers to 1 decimal to match your annot fmt=".1f"
        for row in ws_heat.iter_rows(min_row=2, min_col=2, max_row=1 + n_rows, max_col=1 + n_cols):
            for cell in row:
                cell.number_format = "0.0"

        # Sheet 4: keep the original plots too, as a visual reference
        ws_plots = wb.create_sheet("Plots")
        timeseries_path = os.path.join(out_dir, "timeseries_plot.png")
        heatmap_path = os.path.join(out_dir, "heatmap_plot.png")
        if os.path.exists(timeseries_path):
            ws_plots.add_image(XLImage(timeseries_path), "A1")
        if os.path.exists(heatmap_path):
            ws_plots.add_image(XLImage(heatmap_path), "A80")

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=simulation_results.xlsx"}
        )
    except HTTPException as he:
        raise he
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"{type(e).__name__}: {e}")

@app.get("/export-excel-das")
async def export_excel_das():
    config = load_config()
    out_dir = os.path.dirname(os.path.abspath(config.data.processed_path))
    das_excel_path = os.path.join(out_dir, 'Konversi_Curah_Hujan_DAS.xlsx')

    # Validasi apakah file sumber ada
    if not os.path.exists(das_excel_path):
        raise HTTPException(404, "Laporan DAS belum tersedia. Jalankan simulasi terlebih dahulu.")


    return FileResponse(
        das_excel_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="Konversi_Curah_Hujan_DAS.xlsx"
    )