import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import get_column_letter
import pandas as pd

def generate_das_report(
    data, 
    cn_value, 
    area_km2, 
    n_trials, 
    output_excel="Konversi_Curah_Hujan_DAS_2025_2030.xlsx",
    output_chart="grafik_konversi_das.png"
):
    """
    Fungsi untuk melakukan simulasi Monte Carlo konversi curah hujan ke debit DAS 
    serta mengekspor hasilnya menjadi grafik dan file Excel.
    """
    
    # 1. Pemrosesan Data Curah Hujan
    df_p = pd.DataFrame(data) if isinstance(data, dict) else data
    months = ['Jan', 'Feb', 'Mar', 'Apr', 'Mei', 'Jun', 'Jul', 'Agu', 'Sep', 'Okt', 'Nov', 'Des']
    days_in_month = [31, 28.25, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]

    # Transformasi data ke format tabel memanjang (long format)
    records = []
    for idx, row in df_p.iterrows():
        yr = int(row.get('Tahun', idx))
        for m_idx, m_name in enumerate(months):
            p_month = row[m_name]
            days = days_in_month[m_idx]
            records.append({
                'Tahun': yr,
                'Bulan': m_name,
                'Hari_dalam_Bulan': days,
                'P_bulanan_mm': p_month,
            })

    df_long = pd.DataFrame(records)

    # 2. Parameter DAS & SCS-CN
    S_mm = (25400 / cn_value) - 254
    Ia_mm = 0.2 * S_mm
    Area_m2 = area_km2 * 1e6

    # 3. Simulasi Monte Carlo & Konversi SCS-CN
    np.random.seed(42)
    results = []

    for idx, row in df_long.iterrows():
        p_m = row['P_bulanan_mm']
        days = row['Hari_dalam_Bulan']
        n_days = int(np.round(days))

        q_month_trials = []
        q_peak_trials = []

        for t in range(n_trials):
            p_days = np.random.exponential(scale=1.0, size=n_days)
            if p_days.sum() > 0:
                p_days = p_days / p_days.sum() * p_m
            else:
                p_days = np.zeros(n_days)

            q_days = np.where(
                p_days > Ia_mm, ((p_days - Ia_mm) ** 2) / (p_days - Ia_mm + S_mm), 0.0
            )
            q_month_trials.append(q_days.sum())
            q_peak_trials.append(q_days.max())

        q_mm_mean = np.mean(q_month_trials)
        q_peak_day_mm = np.mean(q_peak_trials)

        vol_m3 = (q_mm_mean / 1000.0) * Area_m2
        q_avg_m3s = vol_m3 / (days * 86400)
        q_peak_m3s = (
            (q_peak_day_mm / 1000.0 * Area_m2) / (8 * 3600)
            if q_peak_day_mm > 0
            else 0
        )

        results.append({
            'Tahun': int(row['Tahun']),
            'Bulan': row['Bulan'],
            'P_mm': round(p_m, 2),
            'CN': cn_value,
            'S_mm': round(S_mm, 2),
            'Ia_mm': round(Ia_mm, 2),
            'Runoff_Q_mm': round(q_mm_mean, 2),
            'Volume_m3': round(vol_m3, 0),
            'Debit_Rerata_m3s': round(q_avg_m3s, 3),
            'Debit_Puncak_m3s': round(q_peak_m3s, 3),
        })

    df_res = pd.DataFrame(results)

    # 4. Pembuatan Grafik (Matplotlib)
    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(12, 10), sharex=True)
    df_res['Periode'] = df_res['Bulan'] + ' ' + df_res['Tahun'].astype(str)
    x = np.arange(len(df_res))

    # Subplot 1: Hujan vs Runoff (mm)
    ax1.bar(x - 0.2, df_res['P_mm'], width=0.4, label='Curah Hujan P (mm)', color='#3498db', alpha=0.85)
    ax1.bar(x + 0.2, df_res['Runoff_Q_mm'], width=0.4, label='Runoff Direct Q (mm)', color='#e74c3c', alpha=0.85)
    ax1.set_ylabel('Kedalaman (mm)', fontsize=11, fontweight='bold')
    ax1.set_title('Konversi Curah Hujan Bulanan Menjadi Runoff (SCS-CN & Monte Carlo)', fontsize=13, fontweight='bold', pad=12)
    ax1.legend(loc='upper right')
    ax1.grid(True, linestyle='--', alpha=0.5)

    # Subplot 2: Hidrograf Debit DAS (m3/s)
    ax2.plot(x, df_res['Debit_Rerata_m3s'], marker='o', color='#2ecc71', linewidth=2, label='Debit Rerata DAS (m³/s)')
    ax2.plot(x, df_res['Debit_Puncak_m3s'], marker='s', color='#e67e22', linewidth=2, linestyle='--', label='Debit Puncak Est. (m³/s)')
    ax2.set_ylabel('Debit DAS (m³/s)', fontsize=11, fontweight='bold')
    ax2.set_title('Hidrograf Debit DAS (Rerata & Puncak Est.)', fontsize=13, fontweight='bold', pad=12)
    ax2.set_xticks(x[::3])
    ax2.set_xticklabels(df_res['Periode'][::3], rotation=45, ha='right')
    ax2.legend(loc='upper right')
    ax2.grid(True, linestyle='--', alpha=0.5)

    plt.tight_layout()
    plt.savefig(output_chart, dpi=300)
    plt.close()

    # 5. Export ke File Excel (.xlsx)
    wb = openpyxl.Workbook()

    # Sheet 1: Ringkasan
    ws1 = wb.active
    ws1.title = 'Ringkasan & Parameter DAS'
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells('A1:G2')
    title_cell = ws1['A1']
    title_cell.value = 'KONVERSI CURAH HUJAN KE DEBIT DAS\nMetode SCS-CN & Simulasi Monte Carlo'
    title_cell.font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws1['A4'] = 'Parameter DAS & SCS-CN'
    ws1['A4'].font = Font(name='Calibri', size=13, bold=True, color='1F4E78')

    params_list = [
        ('Luas DAS (A)', area_km2, 'km²'),
        ('Curve Number (CN)', cn_value, '-'),
        ('Potensi Retensi Maksimum (S)', round(S_mm, 2), 'mm  [(25400/CN) - 254]'),
        ('Abstraksi Awal (Ia)', round(Ia_mm, 2), 'mm  [0.2 * S]'),
        ('Simulasi Monte Carlo', n_trials, 'Iterasi Sebaran Harian / Bulan'),
    ]

    headers_p = ['Parameter', 'Nilai', 'Satuan / Keterangan']
    for col_num, h in enumerate(headers_p, 1):
        cell = ws1.cell(row=5, column=col_num)
        cell.value = h
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9'),
    )

    for row_idx, (p, v, u) in enumerate(params_list, 6):
        ws1.cell(row=row_idx, column=1, value=p).font = Font(name='Calibri', size=11)
        ws1.cell(row=row_idx, column=2, value=v).font = Font(name='Calibri', size=11, bold=True)
        ws1.cell(row=row_idx, column=3, value=u).font = Font(name='Calibri', size=11, italic=True)
        for c in range(1, 4):
            ws1.cell(row=row_idx, column=c).border = thin_border

    ws1['A12'] = 'Ringkasan Tahunan Curah Hujan vs Debit DAS'
    ws1['A12'].font = Font(name='Calibri', size=13, bold=True, color='1F4E78')

    yearly_df = df_res.groupby('Tahun').agg({
        'P_mm': 'sum',
        'Runoff_Q_mm': 'sum',
        'Volume_m3': 'sum',
        'Debit_Rerata_m3s': 'mean',
        'Debit_Puncak_m3s': 'max',
    }).reset_index()

    headers_y = ['Tahun', 'Total CH (mm)', 'Total Runoff (mm)', 'Total Volume (m³)', 'Debit Rerata (m³/s)', 'Debit Puncak Maks (m³/s)']
    for col_num, h in enumerate(headers_y, 1):
        cell = ws1.cell(row=13, column=col_num)
        cell.value = h
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r_idx, r_data in yearly_df.iterrows():
        row_num = 14 + r_idx
        ws1.cell(row=row_num, column=1, value=int(r_data['Tahun'])).alignment = Alignment(horizontal='center')
        ws1.cell(row=row_num, column=2, value=r_data['P_mm']).number_format = '#,##0.0'
        ws1.cell(row=row_num, column=3, value=r_data['Runoff_Q_mm']).number_format = '#,##0.0'
        ws1.cell(row=row_num, column=4, value=r_data['Volume_m3']).number_format = '#,##0'
        ws1.cell(row=row_num, column=5, value=r_data['Debit_Rerata_m3s']).number_format = '0.000'
        ws1.cell(row=row_num, column=6, value=r_data['Debit_Puncak_m3s']).number_format = '0.000'

        fill_color = 'F2F2F2' if r_idx % 2 == 1 else 'FFFFFF'
        for c in range(1, 7):
            cell = ws1.cell(row=row_num, column=c)
            cell.border = thin_border
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')

    # Sheet 2: Tabulasi Bulanan
    ws2 = wb.create_sheet(title='Tabulasi Bulanan')
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells('A1:J1')
    t2 = ws2['A1']
    t2.value = 'HASIL KONVERSI CURAH HUJAN KE RUNOFF & DEBIT DAS'
    t2.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    t2.fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    t2.alignment = Alignment(horizontal='center', vertical='center')

    headers_m = ['Tahun', 'Bulan', 'Curah Hujan P (mm)', 'Curve Number (CN)', 'Retensi S (mm)', 'Abstraksi Ia (mm)', 'Runoff Q (mm)', 'Volume Runoff (m³)', 'Debit Rerata (m³/s)', 'Debit Puncak Est. (m³/s)']

    for col_num, h in enumerate(headers_m, 1):
        cell = ws2.cell(row=3, column=col_num)
        cell.value = h
        cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for r_idx, r_data in df_res.iterrows():
        row_num = 4 + r_idx
        ws2.cell(row=row_num, column=1, value=int(r_data['Tahun'])).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_num, column=2, value=r_data['Bulan']).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_num, column=3, value=r_data['P_mm']).number_format = '#,##0.00'
        ws2.cell(row=row_num, column=4, value=r_data['CN']).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_num, column=5, value=r_data['S_mm']).number_format = '#,##0.00'
        ws2.cell(row=row_num, column=6, value=r_data['Ia_mm']).number_format = '#,##0.00'
        ws2.cell(row=row_num, column=7, value=r_data['Runoff_Q_mm']).number_format = '#,##0.00'
        ws2.cell(row=row_num, column=8, value=r_data['Volume_m3']).number_format = '#,##0'
        ws2.cell(row=row_num, column=9, value=r_data['Debit_Rerata_m3s']).number_format = '0.000'
        ws2.cell(row=row_num, column=10, value=r_data['Debit_Puncak_m3s']).number_format = '0.000'

        fill_color = 'F9FBFD' if (r_idx // 12) % 2 == 1 else 'FFFFFF'
        for c in range(1, 11):
            cell = ws2.cell(row=row_num, column=c)
            cell.border = thin_border
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')

    # Penyesuaian Lebar Kolom Otomatis
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['C'].width = 32

    # Simpan Workbook
    wb.save(output_excel)
    print(f'File Excel berhasil dibuat: {output_excel}')

    return {
        'excel_path': output_excel,
        'chart_path': output_chart,
        'dataframe': df_res,
        'das_parameters': {
            'luas_das_km2': area_km2,
            'cn_value': cn_value,
            'potensi_retensi_maks_s_mm': round(S_mm, 2),
            'abstraksi_awal_ia_mm': round(Ia_mm, 2),
            'iterasi_monte_carlo': n_trials
        }
    }